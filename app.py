import re
import html
from datetime import timezone
from email.utils import parsedate_to_datetime
from collections import Counter

import feedparser
import requests
import streamlit as st

st.set_page_config(
    page_title="IR World News Assistant",
    page_icon="🌍",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---- Styling ----
st.markdown("""
<style>
.stApp { background: #f4f7fb; }
.block-container { padding-top: 1.4rem; padding-bottom: 3rem; max-width: 1440px; }
[data-testid="stSidebar"] { background: linear-gradient(180deg, #102744 0%, #173b60 100%); }
[data-testid="stSidebar"] * { color: #f4f8ff !important; }
.hero {
  background: linear-gradient(110deg, #102744 0%, #14568a 100%);
  color: white; padding: 1.8rem 2rem; border-radius: 18px; margin-bottom: 1.2rem;
  box-shadow: 0 10px 28px rgba(16,39,68,.12);
}
.hero h1 { color: white; font-size: 2.15rem; margin: 0 0 .4rem 0; }
.hero p { color: #e0edfb; font-size: 1rem; margin: 0; }
.stat { background: white; border: 1px solid #e1eaf4; padding: 1rem; border-radius: 14px; min-height: 85px; }
.stat-label { color: #60758d; font-size: .85rem; }
.stat-value { color: #102744; font-size: 1.3rem; font-weight: 750; margin-top: .2rem; }
.news-card { background: white; border: 1px solid #e1eaf4; border-radius: 15px; padding: 1rem 1.2rem; margin: .7rem 0; }
.news-title { font-size: 1.05rem; font-weight: 700; line-height: 1.45; color: #132d4b; margin-bottom: .45rem; }
.news-meta { color: #627991; font-size: .82rem; margin-bottom: .6rem; }
.news-summary { color: #31465d; line-height: 1.6; }
.tag { display: inline-block; background: #eaf3ff; color: #145a96; padding: .2rem .55rem; border-radius: 999px; font-size: .75rem; font-weight: 650; margin-bottom: .45rem; }
.study-card { background: white; border: 1px solid #e1eaf4; padding: 1rem 1.15rem; border-radius: 14px; margin: .5rem 0; }
div.stButton > button { border-radius: 10px; font-weight: 650; min-height: 2.55rem; }
</style>
""", unsafe_allow_html=True)

SOURCES = {
    "BBC News": "https://feeds.bbci.co.uk/news/world/rss.xml",
    "CNN": "https://rss.cnn.com/rss/edition_world.rss",
}
TOPICS = {
    "All international news": [],
    "War & conflict": ["war", "conflict", "military", "peace", "attack", "ceasefire", "missile"],
    "Diplomacy": ["diplomacy", "diplomatic", "foreign policy", "summit", "ambassador", "talks"],
    "United Nations": ["united nations", "security council", "un secretary", "un chief"],
    "China–US relations": ["china", "chinese", "united states", "beijing", "washington", "trade war"],
    "Mongolia foreign policy": ["mongolia", "mongolian", "ulaanbaatar"],
    "Economy & trade": ["economy", "trade", "tariff", "sanctions", "market", "export"],
    "Climate & environment": ["climate", "environment", "emissions", "renewable", "warming"],
}
THEORY_NOTES = {
    "Realism": "Focus on state power, national interests, military capability, security, and competition.",
    "Liberalism": "Focus on cooperation, international organizations, trade, diplomacy, and shared rules.",
    "Constructivism": "Focus on identity, norms, ideas, historical memory, and how actors understand events.",
}
STOPWORDS = set("""
a an the and or but if then than to of in on at for from by with about into over after before
is are was were be been being it its this that these those as their they them we our you your
can could may might will would should have has had do does did not no so such very also more
most some many much other another each both all any who what when where why how which while
international relations news says said new latest according report reports amid during against
""".split())

def clean_text(raw):
    raw = re.sub(r"<[^>]+>", " ", raw or "")
    return re.sub(r"\s+", " ", html.unescape(raw)).strip()

@st.cache_data(ttl=600, show_spinner=False)
def load_feed(source_url):
    response = requests.get(
        source_url,
        headers={"User-Agent": "Mozilla/5.0 (compatible; IRWorldNewsAssistant/1.0)"},
        timeout=15,
    )
    response.raise_for_status()
    return feedparser.parse(response.content).entries

def parse_date(entry):
    raw = entry.get("published") or entry.get("updated") or ""
    if not raw:
        return "Date unavailable", 0
    try:
        dt = parsedate_to_datetime(raw)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.strftime("%d %b %Y"), int(dt.timestamp())
    except (TypeError, ValueError, OverflowError):
        return raw, 0

def get_articles():
    articles, errors = [], []
    for source, url in SOURCES.items():
        try:
            for entry in load_feed(url):
                title = clean_text(entry.get("title", "Untitled"))
                summary = clean_text(entry.get("summary", entry.get("description", "")))
                date_label, timestamp = parse_date(entry)
                articles.append({
                    "source": source, "title": title, "summary": summary,
                    "url": entry.get("link", ""), "date": date_label, "timestamp": timestamp,
                })
        except Exception:
            errors.append(source)
    unique, seen = [], set()
    for article in sorted(articles, key=lambda a: a["timestamp"], reverse=True):
        key = article["title"].casefold()
        if key and key not in seen:
            unique.append(article)
            seen.add(key)
    return unique, errors

# ---- Lecture-note tools: offline / no API key ----
def extract_uploaded_text(uploaded_file):
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("PDF уншихад pypdf сан хэрэгтэй. requirements.txt файлдаа pypdf нэмнэ үү.")
        import io
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("PDF-ээс текст уншиж чадсангүй. Энэ нь зураг хэлбэрийн PDF байж магадгүй.")
        return text
    if name.endswith(".docx"):
        try:
            from docx import Document
        except ImportError:
            raise ValueError("Word уншихад python-docx сан хэрэгтэй. requirements.txt файлдаа python-docx нэмнэ үү.")
        import io
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs)
    raise ValueError("Зөвхөн PDF эсвэл DOCX файл оруулна уу. Хуучин .doc формат дэмжихгүй.")

def split_sentences(text):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9“\"(])|(?<=。)\s*", text)
    return [p.strip(" •-\t") for p in parts if len(p.strip()) > 35]

def extract_keywords(text, limit=12):
    words = re.findall(r"\b[a-zA-Z][a-zA-Z'-]{2,}\b", text.lower())
    counts = Counter(w for w in words if w not in STOPWORDS)
    return [w for w, _ in counts.most_common(limit)]

def make_notes(text, style="Study notes"):
    # Transparent extractive baseline: no cloud AI or API key.
    lines = [re.sub(r"^\s*(?:[-*•]+|\d+[.)])\s*", "", line).strip()
             for line in text.splitlines() if line.strip()]
    sentences = split_sentences(text)
    keywords = extract_keywords(text, 14)
    if not sentences:
        sentences = [line for line in lines if len(line) > 25]
    if not sentences:
        return {"summary": "Оруулах текст хангалтгүй байна. Илүү урт лекцийн текст оруулна уу.",
                "points": [], "terms": keywords, "questions": []}

    # Score sentences for lecture-topic relevance, while preserving original wording.
    scored = []
    for idx, sentence in enumerate(sentences):
        lower = sentence.lower()
        score = sum(2 for key in keywords[:10] if key in lower)
        if any(marker in lower for marker in ["is defined as", "refers to", "because", "therefore", "however",
                                               "main reason", "consequence", "result", "theory", "ideology",
                                               "important", "caused by", "leads to"]):
            score += 2
        if len(sentence) > 300:
            score -= 1
        scored.append((score, idx, sentence))
    top = sorted(scored, key=lambda item: (item[0], -item[1]), reverse=True)
    chosen = sorted(top[:min(8, len(top))], key=lambda item: item[1])
    points = [sentence for _, _, sentence in chosen]
    summary = " ".join(points[:3])
    questions = []
    for term in keywords[:5]:
        questions.append(f"What does “{term}” mean in the context of this lecture?")
    return {"summary": summary, "points": points, "terms": keywords, "questions": questions}

def build_example(notes_text, theory):
    text = notes_text.strip()
    low = text.lower()
    examples = {
        "Realism": (
            "Possible example: competition between the United States and China.",
            "Use this only if your notes discuss state power, security, national interests, or competition. "
            "Explain each country's interests, then connect those interests to the theory's focus on power and security."
        ),
        "Liberalism": (
            "Possible example: countries cooperating through the United Nations.",
            "Use this only if your notes discuss cooperation, institutions, diplomacy, trade, or shared rules. "
            "Explain how an organization or agreement helps states work together."
        ),
        "Constructivism": (
            "Possible example: how national identity and historical memory influence foreign policy.",
            "Use this only if your notes discuss identity, ideas, norms, values, or shared beliefs. "
            "Explain how these ideas shape an actor's understanding of its interests."
        ),
    }
    headline, explanation = examples[theory]
    if theory == "Realism" and not any(w in low for w in ["power", "security", "state", "interest", "realism"]):
        caveat = "Your notes do not clearly mention the usual Realism keywords, so check that this example fits the lecture."
    elif theory == "Liberalism" and not any(w in low for w in ["cooperation", "institution", "trade", "liberal", "organization"]):
        caveat = "Your notes do not clearly mention the usual Liberalism keywords, so check that this example fits the lecture."
    elif theory == "Constructivism" and not any(w in low for w in ["identity", "norm", "idea", "constructiv", "belief", "value"]):
        caveat = "Your notes do not clearly mention the usual Constructivism keywords, so check that this example fits the lecture."
    else:
        caveat = "This is an illustrative example, not a claim that the lecture explicitly used this case."
    return headline, explanation, caveat

# ---- Navigation ----
with st.sidebar:
    st.markdown("## 🌍 IR World News")
    st.caption("News + lecture study workspace")
    page = st.radio(
        "NAVIGATION",
        ["🏠 Home", "📰 Latest News", "📝 Lecture Notes", "📚 IR Theory Lab", "🌏 Mongolia Focus", "ℹ️ About"],
        label_visibility="collapsed",
    )
    st.markdown("---")
    st.markdown("**QUICK LINKS**")
    st.markdown("[BBC World News ↗](https://www.bbc.com/news/world)")
    st.markdown("[CNN World ↗](https://edition.cnn.com/world)")
    st.markdown("[United Nations ↗](https://news.un.org/)")
    st.markdown("---")
    st.caption("For International Relations students")

st.markdown("""
<div class="hero">
  <h1>Stay informed. Think globally. 🌐</h1>
  <p>Follow international headlines, turn lectures into study notes, and practise applying IR theories.</p>
</div>
""", unsafe_allow_html=True)

if "articles" not in st.session_state:
    st.session_state.articles = []
if "feed_errors" not in st.session_state:
    st.session_state.feed_errors = []

if page == "📝 Lecture Notes":
    st.markdown("## 📝 Lecture Notes Studio")
    st.write("PDF, DOCX эсвэл лекцийн текстээ оруулаад тэмдэглэл, гол нэр томьёо, жишээ, шалгалтын асуулт бэлдээрэй.")
    st.info("API түлхүүргүй хувилбар: таны оруулсан текстийг төхөөрөмж дээр биш, Streamlit апп ажиллаж буй орчинд боловсруулна. Энэ нь сонгосон өгүүлбэрүүдийг ялгаж, түлхүүр үг олдог энгийн арга бөгөөд бүрэн AI ойлголт биш.")
    tab_file, tab_text = st.tabs(["📄 Upload PDF / Word", "✍️ Paste lecture text"])
    with tab_file:
        uploaded = st.file_uploader("Choose a lecture file", type=["pdf", "docx"], help="PDF эсвэл .docx файл. Хуучин .doc формат дэмжихгүй.")
        file_text = ""
        if uploaded:
            try:
                file_text = extract_uploaded_text(uploaded)
                st.success(f"Текст уншлаа: ойролцоогоор {len(file_text.split())} үг.")
            except Exception as exc:
                st.error(str(exc))
    with tab_text:
        pasted_text = st.text_area("Paste your lecture transcript or notes here", height=260,
                                   placeholder="Paste lecture transcript, textbook notes, or class notes…")
    input_mode = st.radio("Тэмдэглэл гаргах эх сурвалж", ["Uploaded file", "Pasted text"], horizontal=True)
    source_text = file_text if input_mode == "Uploaded file" else pasted_text
    note_style = st.selectbox("Output style", ["Study notes", "Short exam revision", "Key concepts and definitions"])
    if st.button("✨ Create lecture notes", type="primary", use_container_width=True):
        if len(source_text.strip()) < 100:
            st.warning("Илүү дэлгэрэнгүй лекц оруулна уу (дор хаяж 100 тэмдэгт).")
        else:
            st.session_state.generated_notes = make_notes(source_text, note_style)
            st.session_state.original_lecture = source_text
            st.session_state.note_style = note_style

    if "generated_notes" in st.session_state:
        notes = st.session_state.generated_notes
        st.markdown("---")
        st.markdown("### Your lecture notes")
        st.markdown("**Main idea / гол санаа**")
        st.write(notes["summary"])
        st.markdown("**Key points / чухал санаанууд**")
        for point in notes["points"]:
            st.markdown(f"- {point}")
        st.markdown("**Keywords / түлхүүр үгс**")
        if notes["terms"]:
            st.write(" · ".join(notes["terms"]))
        else:
            st.write("Түлхүүр үг илэрсэнгүй.")
        st.markdown("### 🧠 Generate an example from these notes")
        theory = st.selectbox("Choose an IR theory", list(THEORY_NOTES.keys()))
        if st.button("Create theory-based example"):
            example, explanation, caveat = build_example(st.session_state.original_lecture, theory)
            st.session_state.last_example = (theory, example, explanation, caveat)
        if "last_example" in st.session_state:
            theory_saved, example, explanation, caveat = st.session_state.last_example
            st.markdown(f'<div class="study-card"><span class="tag">{html.escape(theory_saved)}</span><h4>{html.escape(example)}</h4><p>{html.escape(explanation)}</p><p><b>Note:</b> {html.escape(caveat)}</p></div>', unsafe_allow_html=True)
        st.markdown("### 📝 Practice questions")
        for i, question in enumerate(notes["questions"], start=1):
            st.markdown(f"{i}. {question}")
        st.download_button("⬇️ Download notes as TXT", data=(
            "IR WORLD NEWS ASSISTANT — LECTURE NOTES\n\nMAIN IDEA\n" + notes["summary"] +
            "\n\nKEY POINTS\n" + "\n".join("- " + p for p in notes["points"]) +
            "\n\nKEYWORDS\n" + ", ".join(notes["terms"]) +
            "\n\nPRACTICE QUESTIONS\n" + "\n".join(notes["questions"])
        ), file_name="ir_lecture_notes.txt", mime="text/plain")
    st.caption("Tip: Review the generated notes against your original lecture. This no-key version may miss context or select less important sentences.")
elif page == "📚 IR Theory Lab":
    st.markdown("## 📚 IR Theory Lab")
    st.write("Use these lenses to analyse a current event.")
    for theory, description in THEORY_NOTES.items():
        with st.expander(theory, expanded=(theory == "Realism")):
            st.write(description)
            if theory == "Realism":
                st.write("Guiding questions: Which actors have power? What security concerns or national interests are involved?")
            elif theory == "Liberalism":
                st.write("Guiding questions: Which institutions, agreements, or shared interests encourage cooperation?")
            else:
                st.write("Guiding questions: How do identities, norms, or historical narratives shape decisions?")
    st.markdown("**Short-answer structure**")
    st.code("1. Define the theory.\\n2. Explain its main idea.\\n3. Give one relevant example.\\n4. Explain why the example fits.")
elif page == "ℹ️ About":
    st.markdown("## About IR World News Assistant")
    st.write("A study dashboard for International Relations students. It combines publisher RSS headlines with an offline-style lecture-note helper.")
    st.warning("The lecture helper does not use a large language model. Its notes and examples are a first draft; verify facts and course concepts against the original lecture.")
elif page == "🌏 Mongolia Focus":
    if not st.session_state.articles:
        with st.spinner("Loading news feeds..."):
            st.session_state.articles, st.session_state.feed_errors = get_articles()
    q = st.text_input("Search Mongolia-related news", placeholder="Mongolia, Ulaanbaatar, foreign policy…")
    visible = [a for a in st.session_state.articles if any(k in (a["title"] + " " + a["summary"]).lower() for k in TOPICS["Mongolia foreign policy"])]
    if q.strip():
        visible = [a for a in visible if q.casefold() in (a["title"] + " " + a["summary"]).casefold()]
    if not visible:
        st.info("No Mongolia-related headlines found in the currently loaded feeds. Try refreshing or search publisher websites.")
    for a in visible[:20]:
        st.markdown(f'<div class="news-card"><div class="tag">MONGOLIA FOCUS</div><div class="news-title">{html.escape(a["title"])}</div><div class="news-meta">{html.escape(a["source"])} · {html.escape(a["date"])}</div><p>{html.escape(a["summary"][:500])}</p><a href="{html.escape(a["url"], quote=True)}" target="_blank" rel="noopener noreferrer">Read original article ↗</a></div>', unsafe_allow_html=True)
else:
    if page == "🏠 Home":
        c1, c2, c3 = st.columns(3)
        with c1:
            st.markdown('<div class="stat"><div class="stat-label">🗞️ News sources</div><div class="stat-value">BBC · CNN</div></div>', unsafe_allow_html=True)
        with c2:
            st.markdown('<div class="stat"><div class="stat-label">📝 Study tools</div><div class="stat-value">Lecture notes</div></div>', unsafe_allow_html=True)
        with c3:
            st.markdown('<div class="stat"><div class="stat-label">🎓 Practice</div><div class="stat-value">IR theories</div></div>', unsafe_allow_html=True)
    if not st.session_state.articles:
        with st.spinner("Connecting to BBC and CNN feeds..."):
            st.session_state.articles, st.session_state.feed_errors = get_articles()
    if st.session_state.feed_errors:
        st.warning("Could not reach: " + ", ".join(st.session_state.feed_errors) + ". A publisher feed may be temporarily unavailable.")
    col_left, col_right = st.columns([4, 1])
    with col_left:
        st.markdown("## 📰 International news desk")
        st.caption("Publisher RSS headlines · Open original stories for full context")
    with col_right:
        if st.button("🔄 Refresh news", type="primary", use_container_width=True):
            load_feed.clear()
            st.session_state.articles, st.session_state.feed_errors = get_articles()
            st.rerun()
    f1, f2, f3 = st.columns([2, 2, 2])
    with f1:
        search = st.text_input("🔎 Search headlines", placeholder="Country, event, keyword…")
    with f2:
        topic = st.selectbox("Topic", list(TOPICS.keys()))
    with f3:
        source_filter = st.selectbox("Source", ["All sources", *SOURCES.keys()])
    visible = st.session_state.articles
    if TOPICS[topic]:
        visible = [a for a in visible if any(k in (a["title"] + " " + a["summary"]).lower() for k in TOPICS[topic])]
    if source_filter != "All sources":
        visible = [a for a in visible if a["source"] == source_filter]
    if search.strip():
        term = search.casefold()
        visible = [a for a in visible if term in (a["title"] + " " + a["summary"]).casefold()]
    st.caption(f"Showing {min(len(visible), 20)} of {len(visible)} matching headlines")
    if not visible:
        st.info("No matching headlines. Try another topic or search term.")
    for a in visible[:20]:
        st.markdown(
            f'<div class="news-card"><div class="tag">{html.escape(a["source"].upper())}</div>'
            f'<div class="news-title">{html.escape(a["title"])}</div>'
            f'<div class="news-meta">🗓️ {html.escape(a["date"])} &nbsp; · &nbsp; 📰 {html.escape(a["source"])}</div>'
            f'<div class="news-summary">{html.escape(a["summary"][:600]) or "No summary supplied by the feed."}</div>'
            f'<br><a href="{html.escape(a["url"], quote=True)}" target="_blank" rel="noopener noreferrer">Read original article ↗</a></div>',
            unsafe_allow_html=True,
        )

st.markdown("---")
with st.expander("🎓 Study tip: connect news to IR theory"):
    chosen_theory = st.selectbox("Choose a theory", list(THEORY_NOTES.keys()), key="footer_theory")
    st.write(THEORY_NOTES[chosen_theory])
st.caption("IR World News Assistant · Educational use · Verify lecture notes and news against original sources.")
